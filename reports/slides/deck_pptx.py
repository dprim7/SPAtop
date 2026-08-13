#!/usr/bin/env python3
"""PPTX back end for the slide standard: same DECK dict, editable output.

`deck.py` renders the deck to self-contained HTML. This module renders the
SAME content dict to a native .pptx so it can be edited in PowerPoint,
Keynote or Google Slides. Geometry, palette, type scale and logo placement
mirror the HTML renderer, so what you see in the browser is what you get in
the deck, with every element a real editable shape (no flattened images
except the figures themselves).

    from deck_pptx import render_pptx
    render_pptx(DECK, "out.pptx")

Everything is a native object:
  * titles, bullets, kickers, callouts  -> text frames
  * tables and minitables               -> real PowerPoint tables
  * equations                           -> text frames with runs for sub/sup
  * figures and logos                   -> pictures
"""
import os
import re
import html as _html

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

import deck as _theme  # palette, asset resolution, LOGO_SPLIT

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

# The HTML slide is 1280 x 720 CSS px; a 16:9 PowerPoint slide is
# 13.333 x 7.5 in. One CSS px maps to exactly 9525 EMU at this scale.
PX = 9525
W_PX, H_PX = 1280, 720


def _rgb(hexstr):
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


NAVY, BLUE, GOLD = _rgb(_theme.NAVY), _rgb(_theme.BLUE), _rgb(_theme.GOLD)
INK, MUTED = _rgb(_theme.INK), _rgb(_theme.MUTED)
WASH, WHITE = _rgb(_theme.WASH), _rgb("#FFFFFF")
RULE = _rgb("#E2E7EC")
SANS = "Helvetica Neue"
MONO = "Consolas"
SERIF = "Georgia"

# layout constants, in CSS px to match deck.py
PAD_X, PAD_TOP = 72, 58
BODY_TOP, BODY_H = 162, 474
FOOT_Y = 660


def _mark_path(kind, dark):
    stem = {"inst": "ucsd-logo", "exp": "cms-logo"}[kind]
    order = [f"{stem}-white", stem] if dark else [stem, f"{stem}-white"]
    for s in order:
        for ext in ("png", "jpg"):
            p = os.path.join(ASSETS, f"{s}.{ext}")
            if os.path.exists(p):
                return p
    return None


def _fig_path(rel):
    if not rel:
        return None
    p = rel if os.path.isabs(rel) else os.path.join(HERE, rel)
    return p if os.path.exists(p) else None


# ---------------------------------------------------------------- rich text

_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|<sub>.+?</sub>|<sup>.+?</sup>|"
                    r"<i>.+?</i>|<span class='v'>.+?</span>)")


def _add_rich(para, text, size, color=INK, font=SANS, bold_color=None):
    """Render the deck's tiny inline vocabulary into runs.

    The markup nests (``<i>H<sub>T</sub></i>``, ``<span class='v'>10<sup>4</sup>
    </span>``), so this walks it recursively, accumulating a style rather than
    matching one tag and stopping. A flat pass leaves literal tags in the text
    and silently drops inner superscripts.
    """
    base = {"bold": False, "italic": False, "mono": False, "serif": False,
            "scale": 1.0, "baseline": None, "color": color}
    _emit(para, _html.unescape(text), base, size, font, bold_color)


def _emit(para, text, st, size, font, bold_color):
    for piece in _TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            _emit(para, piece[2:-2],
                  {**st, "bold": True, "color": bold_color or NAVY},
                  size, font, bold_color)
        elif piece.startswith("`"):
            _emit(para, piece[1:-1],
                  {**st, "mono": True, "scale": st["scale"] * 0.92},
                  size, font, bold_color)
        elif piece.startswith("<sub>"):
            _emit(para, piece[5:-6],
                  {**st, "scale": st["scale"] * 0.62, "baseline": "-25000"},
                  size, font, bold_color)
        elif piece.startswith("<sup>"):
            _emit(para, piece[5:-6],
                  {**st, "scale": st["scale"] * 0.62, "baseline": "30000"},
                  size, font, bold_color)
        elif piece.startswith("<i>"):
            _emit(para, piece[3:-4], {**st, "italic": True, "serif": True},
                  size, font, bold_color)
        elif piece.startswith("<span"):
            _emit(para, piece[len("<span class='v'>"):-len("</span>")],
                  {**st, "bold": True, "color": BLUE}, size, font, bold_color)
        else:
            plain = re.sub(r"<[^>]+>", "", piece)
            if not plain:
                continue
            run = para.add_run()
            run.text = plain
            f = run.font
            f.size = Pt(size * st["scale"])
            f.name = MONO if st["mono"] else (SERIF if st["serif"] else font)
            f.color.rgb = st["color"]
            f.bold = st["bold"]
            f.italic = st["italic"]
            if st["baseline"]:
                f._rPr.set("baseline", st["baseline"])


def _box(slide, x, y, w, h):
    return slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX),
                                    Emu(w * PX), Emu(h * PX))


def _rect(slide, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x * PX), Emu(y * PX),
                                Emu(w * PX), Emu(h * PX))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _pic(slide, path, x, y, max_w, max_h):
    """Place an image scaled to fit the box, centred in it."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    return slide.shapes.add_picture(
        path, Emu(int((x + (max_w - w) / 2) * PX)),
        Emu(int((y + (max_h - h) / 2) * PX)), Emu(int(w * PX)), Emu(int(h * PX)))


def _marks(slide, dark=False, big=False, y=None):
    """Logo pair, matching the HTML: CMS then UCSD, right aligned."""
    cms_h = 124 if big else 60
    ucsd_h = 68 if big else 38
    gap = 14
    paths = [(_mark_path("exp", dark), cms_h), (_mark_path("inst", dark), ucsd_h)]
    widths = []
    from PIL import Image
    for p, hh in paths:
        if p:
            with Image.open(p) as im:
                widths.append(im.size[0] * hh / im.size[1])
        else:
            widths.append(0)
    total = sum(widths) + gap * (len([p for p, _ in paths if p]) - 1)
    x = W_PX - (92 if big else PAD_X) - total
    top = y if y is not None else (44 if big else 56)
    for (p, hh), w in zip(paths, widths):
        if not p:
            continue
        slide.shapes.add_picture(p, Emu(int(x * PX)),
                                 Emu(int((top + (cms_h - hh) / 2) * PX)),
                                 Emu(int(w * PX)), Emu(int(hh * PX)))
        x += w + gap


def _title_slide(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(sl, 0, 0, W_PX, H_PX, NAVY)
    _marks(sl, dark=True, big=True, y=44)
    tb = _box(sl, 92, 236, 1020, 130)
    tf = tb.text_frame
    tf.word_wrap = True
    _add_rich(tf.paragraphs[0], s["title"], 44, WHITE, SANS, bold_color=WHITE)
    tf.paragraphs[0].runs[0].font.bold = True
    _rect(sl, 92, 392, 132, 6, GOLD)
    if s.get("subtitle"):
        sb = _box(sl, 92, 416, 940, 90)
        sb.text_frame.word_wrap = True
        _add_rich(sb.text_frame.paragraphs[0], s["subtitle"], 20,
                  _rgb("#C8D4E2"), SANS, bold_color=WHITE)
    if s.get("meta"):
        mb = _box(sl, 92, 600, 1000, 40)
        _add_rich(mb.text_frame.paragraphs[0],
                  "   ·   ".join(s["meta"]), 13, _rgb("#8FA3BA"), SANS)
    return sl


def _divider_slide(prs, s):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(sl, 0, 0, W_PX, H_PX, NAVY)
    _marks(sl, dark=True, big=True, y=44)
    tb = _box(sl, 0, 300, W_PX, 90)
    tf = tb.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _add_rich(tf.paragraphs[0], s["title"], 47, WHITE, SANS, bold_color=WHITE)
    tf.paragraphs[0].runs[0].font.bold = True
    _rect(sl, (W_PX - 132) / 2, 410, 132, 6, GOLD)
    return sl


def _header(sl, title, num, footer):
    _marks(sl, dark=False, big=False, y=26)
    tb = _box(sl, PAD_X, PAD_TOP - 10, 820, 96)
    tf = tb.text_frame
    tf.word_wrap = True
    _add_rich(tf.paragraphs[0], title, 29, NAVY, SANS, bold_color=NAVY)
    for r in tf.paragraphs[0].runs:
        r.font.bold = True
    _rect(sl, PAD_X, 140, 104, 5, GOLD)
    _rect(sl, PAD_X, FOOT_Y - 12, W_PX - 2 * PAD_X, 1, RULE)
    fb = _box(sl, PAD_X, FOOT_Y, 900, 30)
    _add_rich(fb.text_frame.paragraphs[0], footer, 11, MUTED, SANS)
    nb = _box(sl, W_PX - PAD_X - 120, FOOT_Y, 120, 30)
    nb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    _add_rich(nb.text_frame.paragraphs[0], str(num), 11, MUTED, SANS)


def _table(sl, tbl, x, y, w, mini=False):
    rows, cols = len(tbl["rows"]) + 1, len(tbl["head"])
    rh = 26 if mini else 38
    h = rh * rows
    shape = sl.shapes.add_table(rows, cols, Emu(int(x * PX)), Emu(int(y * PX)),
                                Emu(int(w * PX)), Emu(int(h * PX)))
    t = shape.table
    t.first_row = True
    for j, htxt in enumerate(tbl["head"]):
        c = t.cell(0, j)
        c.text = ""
        p = c.text_frame.paragraphs[0]
        _add_rich(p, htxt, 11 if mini else 12,
                  BLUE if mini else WHITE, SANS)
        for r_ in p.runs:
            r_.font.bold = True
        c.fill.solid()
        c.fill.fore_color.rgb = WASH if mini else NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(tbl["rows"], start=1):
        cells = row["cells"] if isinstance(row, dict) else row
        win = isinstance(row, dict) and row.get("win")
        for j, val in enumerate(cells):
            c = t.cell(i, j)
            c.text = ""
            p = c.text_frame.paragraphs[0]
            _add_rich(p, str(val), 12 if mini else 15,
                      NAVY if (j == 0 or win) else INK,
                      SANS if j == 0 else MONO)
            if win:
                for r_ in p.runs:
                    r_.font.bold = True
            c.fill.solid()
            c.fill.fore_color.rgb = (GOLD if win else
                                     (WASH if (i % 2 == 0 and not mini) else WHITE))
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return h


def _content_slide(prs, s, num, footer):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _header(sl, s["title"], num, footer)
    fig = _fig_path(s.get("figure"))
    has_bul = bool(s.get("bullets"))
    y = BODY_TOP
    left_w = W_PX - 2 * PAD_X

    if fig and has_bul:                       # split layout
        fw = int(left_w * 0.56)
        _pic(sl, fig, PAD_X, y, fw, BODY_H - 20)
        bx, bw = PAD_X + fw + 44, left_w - fw - 44
        _bullets(sl, s["bullets"], bx, y + 10, bw)
        if s.get("caption"):
            cb = _box(sl, PAD_X, y + BODY_H - 26, fw, 30)
            cb.text_frame.word_wrap = True
            _add_rich(cb.text_frame.paragraphs[0], s["caption"], 10, MUTED, SANS)
        return sl

    if fig:
        cap_h = 34 if s.get("caption") else 0
        _pic(sl, fig, PAD_X, y, left_w, BODY_H - cap_h)
        if s.get("caption"):
            cb = _box(sl, PAD_X, y + BODY_H - cap_h + 2, left_w, 30)
            cb.text_frame.word_wrap = True
            p = cb.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _add_rich(p, s["caption"], 10, MUTED, SANS)
        return sl

    if s.get("table"):
        y += _table(sl, s["table"], PAD_X, y, left_w) + 22
    if has_bul:
        y = _bullets(sl, s["bullets"], PAD_X, y, left_w)
    if s.get("kicker"):
        kb = _box(sl, PAD_X, y + 4, left_w, 40)
        kb.text_frame.word_wrap = True
        p = kb.text_frame.paragraphs[0]
        _add_rich(p, s["kicker"], 17, BLUE, SANS, bold_color=BLUE)
        for r_ in p.runs:
            r_.font.bold = True
        y += 44
    if s.get("equation"):
        eqs = s["equation"]
        eqs = eqs if isinstance(eqs, list) else [eqs]
        for e in eqs:
            _rect(sl, PAD_X, y, left_w, 54, WASH)
            eb = _box(sl, PAD_X + 22, y + 10, left_w - 44, 40)
            p = eb.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _add_rich(p, e, 19, NAVY, SERIF, bold_color=BLUE)
            y += 66
    if s.get("minitable"):
        mt = s["minitable"]
        rows = len(mt["rows"]) + 1
        panel_h = 26 * rows + (30 if mt.get("label") else 10) + 10
        _rect(sl, PAD_X, y, left_w, panel_h, _rgb("#EEF4F9"))
        _rect(sl, PAD_X, y, 5, panel_h, BLUE)
        yy = y + 8
        if mt.get("label"):
            lb = _box(sl, PAD_X + 22, yy, left_w - 44, 26)
            p = lb.text_frame.paragraphs[0]
            _add_rich(p, mt["label"].upper(), 10, BLUE, SANS)
            for r_ in p.runs:
                r_.font.bold = True
            yy += 24
        _table(sl, {"head": mt["head"], "rows": mt["rows"]},
               PAD_X + 22, yy, left_w - 44, mini=True)
        y += panel_h + 12
    if s.get("callout"):
        c = s["callout"]
        c = c if isinstance(c, dict) else {"text": c}
        ch = 92 if c.get("label") else 66
        cy = BODY_TOP + BODY_H - ch
        _rect(sl, PAD_X, cy, left_w, ch, _rgb("#EEF4F9"))
        _rect(sl, PAD_X, cy, 5, ch, BLUE)
        yy = cy + 12
        if c.get("label"):
            lb = _box(sl, PAD_X + 22, yy, left_w - 44, 24)
            p = lb.text_frame.paragraphs[0]
            _add_rich(p, c["label"].upper(), 10, BLUE, SANS)
            for r_ in p.runs:
                r_.font.bold = True
            yy += 24
        tb = _box(sl, PAD_X + 22, yy, left_w - 44, ch - (yy - cy) - 8)
        tb.text_frame.word_wrap = True
        _add_rich(tb.text_frame.paragraphs[0], c["text"], 15, INK, SANS)
        y = max(y, cy + ch)
    if y > BODY_TOP + BODY_H + 1:
        _OVERFLOW.append((num, round(y)))
    return sl


def _bullets(sl, items, x, y, w):
    for b in items:
        _rect(sl, x, y + 13, 11, 3, GOLD)
        tb = _box(sl, x + 26, y - 4, w - 26, 60)
        tf = tb.text_frame
        tf.word_wrap = True
        _add_rich(tf.paragraphs[0], b, 18, INK, SANS)
        y += 46
    return y


def _decisions_slide(prs, s, num, footer):
    from pptx.enum.shapes import MSO_SHAPE
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _header(sl, s["title"], num, footer)
    y = BODY_TOP
    for i, d in enumerate(s["decisions"], start=1):
        c = sl.shapes.add_shape(MSO_SHAPE.OVAL, Emu(PAD_X * PX), Emu(int(y * PX)),
                                Emu(34 * PX), Emu(34 * PX))
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.line.fill.background()
        c.shadow.inherit = False
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _add_rich(p, str(i), 13, WHITE, SANS, bold_color=WHITE)
        for r_ in p.runs:
            r_.font.bold = True
        tb = _box(sl, PAD_X + 52, y - 2, W_PX - 2 * PAD_X - 52, 60)
        tb.text_frame.word_wrap = True
        _add_rich(tb.text_frame.paragraphs[0], d, 17, INK, SANS)
        y += 54
    return sl


_OVERFLOW = []   # (slide number, body bottom in px) recorded during layout


def render_pptx(deck, out_path):
    _OVERFLOW.clear()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W_PX * PX), Emu(H_PX * PX)
    footer = deck.get("footer", "")
    n = 0
    for s in deck["slides"]:
        t = s.get("type", "content")
        if t == "title":
            _title_slide(prs, s)
        elif t == "divider":
            _divider_slide(prs, s)
        elif t == "decisions":
            n += 1
            _decisions_slide(prs, s, n, footer)
        else:
            n += 1
            _content_slide(prs, s, n, footer)
    if _OVERFLOW:
        print("WARNING: body content overruns the footer rule on slide(s) "
              + ", ".join(f"{i} (bottom {b}px, limit {BODY_TOP + BODY_H})"
                          for i, b in _OVERFLOW))
    prs.save(out_path)
    return out_path
